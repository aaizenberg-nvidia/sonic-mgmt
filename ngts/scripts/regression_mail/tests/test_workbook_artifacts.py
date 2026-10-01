from __future__ import annotations

import hashlib
import json
import stat
import tempfile
import unittest
from pathlib import Path

import openpyxl

from ngts.scripts.regression_mail.grouping import build_deterministic_report
from ngts.scripts.regression_mail.models import SemanticGroup
from ngts.scripts.regression_mail.redaction import SetupRedactor
from ngts.scripts.regression_mail.workbook import (
    ArtifactWriter,
    FALLBACK_SKIP_REASON,
    REQUIRED_HEADERS,
    WorkbookLoader,
    _display_testbed,
)


VERSION = "202608_RC.17-deadbeef_Internal"


class WorkbookArtifactTest(unittest.TestCase):
    def test_explicit_unknown_platform_falls_back_to_testbed(self) -> None:
        self.assertEqual(
            "tb1",
            _display_testbed("Unknown platform", "t0", "tb1"),
        )

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / "results.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "community_tests"
        headers = list(REQUIRED_HEADERS) + ["extra", "platform"]
        sheet.append(headers)
        sheet.append(
            [
                "s1",
                "k1",
                "tb1",
                "tests/a.py::test_x[a]",
                "fail",
                "failure text",
                "t0",
                "sku",
                "SONiC." + VERSION,
                "tests/a.py::test_x[a]",
                "keep",
                "SN5600",
            ]
        )
        sheet.append(
            [
                "s2",
                "k2",
                "tb2",
                "tests/b.py::test_skip",
                "skipped",
                "Internal ticket",
                "t1",
                "sku2",
                VERSION,
                "tests/b.py::test_skip",
                "keep2",
                "SN5600",
            ]
        )
        sheet.append(
            [
                "s3",
                "k3",
                "tb3",
                "tests/c.py::test_other",
                "pass",
                "",
                "t2",
                "sku3",
                "202605_RC.1-other_Internal",
                "tests/c.py::test_other",
                "other-version",
                "SN2700",
            ]
        )
        workbook.create_sheet("metadata")["A1"] = "preserve"
        workbook.save(self.path)
        workbook.close()

    def test_filters_exact_version_and_preserves_source(self) -> None:
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        snapshot = WorkbookLoader().load(self.path, VERSION)
        after = hashlib.sha256(self.path.read_bytes()).hexdigest()

        self.assertEqual(before, after)
        self.assertEqual("community_tests", snapshot.sheet_name)
        self.assertEqual(3, len(snapshot.selected_rows))
        self.assertEqual({"fail": 1, "pass": 1, "skipped": 1}, snapshot.result_counts)
        self.assertEqual(
            [("sku", "t0"), ("sku2", "t1"), ("sku3", "t2")],
            snapshot.hardware_pairs,
        )
        self.assertEqual(
            ["202605_RC.1-other_Internal", VERSION],
            snapshot.versions,
        )
        self.assertEqual([], snapshot.issue_skipped)
        self.assertIsNone(snapshot.coverage)
        self.assertEqual(50.0, snapshot.pass_rate)

    def test_atomic_skip_mapping(self) -> None:
        snapshot = WorkbookLoader().load(self.path, VERSION)
        semantic = build_deterministic_report(snapshot)
        skip_group = semantic.skip_groups[0]
        semantic.skip_groups[0] = SemanticGroup(
            group_id=skip_group.group_id,
            member_ids=skip_group.member_ids,
            test_display=skip_group.test_display,
            testbeds=skip_group.testbeds,
            comments="",
            internal_comments="https://redmine.mellanox.com/issues/123",
            redmine_urls=["https://redmine.mellanox.com/issues/123"],
        )
        writer = ArtifactWriter()
        skips_path = writer.write_skips(snapshot, semantic)

        self.assertEqual(
            {
                "https://redmine.mellanox.com/issues/123":
                    FALLBACK_SKIP_REASON
            },
            json.loads(skips_path.read_text(encoding="utf-8")),
        )
        self.assertEqual(0o664, stat.S_IMODE(skips_path.stat().st_mode))

    def test_remove_failed_testcase_rows_are_excluded_before_metrics(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["community_tests"]
        headers = [cell.value for cell in sheet[1]]
        if "modify_xml" not in headers:
            modify_column = len(headers) + 1
            sheet.cell(row=1, column=modify_column, value="modify_xml")
        else:
            modify_column = headers.index("modify_xml") + 1
        sheet.cell(row=2, column=modify_column, value="REMOVE_FAILED_TESTCASE")
        workbook.save(self.path)
        workbook.close()

        snapshot = WorkbookLoader().load(self.path, VERSION)

        self.assertEqual({"fail": 0, "pass": 1, "skipped": 1}, snapshot.result_counts)
        self.assertFalse(any(row.mars_key_id == "k1" for row in snapshot.selected_rows))

    def test_setup_names_are_replaced_in_semantic(self) -> None:
        snapshot = WorkbookLoader().load(self.path, VERSION)
        snapshot.row_comments["s1"] = "analysis on tb1"
        semantic = build_deterministic_report(snapshot)
        semantic.failure_groups[0].comments = "summary on tb1"
        semantic.failure_groups[0].internal_comments = "summary on tb1"
        semantic.failure_groups[0].internal_review_warning = True
        redactor = SetupRedactor(
            snapshot.setup_platforms,
            snapshot.setup_testbeds,
        )
        redactor.redact_semantic(semantic)

        self.assertEqual(
            ["SN5600-t0"],
            semantic.failure_groups[0].testbeds,
        )
        self.assertNotIn("tb1", semantic.failure_groups[0].comments)
        self.assertNotIn("tb1", semantic.failure_groups[0].internal_comments)

    def test_semantically_duplicate_comment_headers_are_rejected(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["community_tests"]
        sheet.cell(1, sheet.max_column + 1).value = "Comments"
        sheet.cell(1, sheet.max_column + 1).value = "comments"
        workbook.save(self.path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "duplicate column names"):
            WorkbookLoader().load(self.path, VERSION)

    def test_setup_mapping_preserves_raw_whitespace_variant_for_redaction(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["community_tests"]
        sheet["C2"] = "tb \n  1"
        sheet["F2"] = "failure on tb \n  1"
        workbook.save(self.path)
        workbook.close()
        snapshot = WorkbookLoader().load(self.path, VERSION)
        semantic = build_deterministic_report(snapshot)
        redactor = SetupRedactor(snapshot.setup_platforms)
        redactor.redact_semantic(semantic)
        group = semantic.failure_groups[0]
        self.assertNotIn("tb \n  1", group.comments)
        self.assertNotIn("tb 1", group.comments)

    def test_ambiguous_or_missing_platform_keeps_testbed(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["community_tests"]
        sheet.append(
            [
                "s4",
                "k4",
                "tb1",
                "tests/d.py::test_other",
                "pass",
                "",
                "t3",
                "sku4",
                "202605_RC.1-other_Internal",
                "tests/d.py::test_other",
                "",
                "SN2201",
            ]
        )
        workbook.save(self.path)
        workbook.close()

        snapshot = WorkbookLoader().load(self.path, VERSION)
        redactor = SetupRedactor(snapshot.setup_platforms)
        self.assertEqual(["tb1"], redactor.platforms(["tb1"]))
        self.assertEqual("", snapshot.setup_testbeds["tb1"])

    def test_missing_skip_reason_uses_contract_fallback(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        workbook["community_tests"]["F3"] = ""
        workbook.save(self.path)
        workbook.close()
        snapshot = WorkbookLoader().load(self.path, VERSION)
        semantic = build_deterministic_report(snapshot)
        semantic.skip_groups[0].redmine_urls = ["https://redmine.mellanox.com/issues/999"]
        path = ArtifactWriter().write_skips(snapshot, semantic)
        self.assertEqual(
            FALLBACK_SKIP_REASON,
            json.loads(path.read_text(encoding="utf-8"))[
                "https://redmine.mellanox.com/issues/999"
            ],
        )

    def test_skip_mapping_uses_one_wip_value_for_every_ticket(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["community_tests"]
        reason_a = "https://redmine.mellanox.com/issues/111 Reason A"
        reason_b = "https://redmine.mellanox.com/issues/222 Reason B"
        sheet["F3"] = reason_a
        sheet.append(
            [
                "s4",
                "k4",
                "tb2",
                "tests/b.py::test_other",
                "skipped",
                reason_b,
                "t1",
                "sku2",
                VERSION,
                "tests/b.py::test_other",
                "",
            ]
        )
        workbook.save(self.path)
        workbook.close()

        snapshot = WorkbookLoader().load(self.path, VERSION)
        semantic = build_deterministic_report(snapshot)
        payload = json.loads(
            ArtifactWriter()
            .write_skips(snapshot, semantic)
            .read_text(encoding="utf-8")
        )

        self.assertEqual(
            FALLBACK_SKIP_REASON,
            payload["https://redmine.mellanox.com/issues/111"],
        )
        self.assertEqual(
            FALLBACK_SKIP_REASON,
            payload["https://redmine.mellanox.com/issues/222"],
        )

    def test_skip_rows_fold_only_matching_reasons(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["community_tests"]
        for key, name, reason in (
            ("k4", "tests/b.py::test_other", "Internal ticket"),
            ("k5", "tests/b.py::test_param[a]", "Other condition"),
            ("k6", "tests/b.py::test_param[b]", "Internal ticket"),
        ):
            sheet.append(
                [
                    key,
                    key,
                    "tb2",
                    name,
                    "skipped",
                    reason,
                    "t1",
                    "sku2",
                    VERSION,
                    name,
                    "",
                ]
            )
        workbook.save(self.path)
        workbook.close()

        semantic = build_deterministic_report(
            WorkbookLoader().load(self.path, VERSION)
        )
        self.assertEqual(2, len(semantic.skip_groups))
        by_comment = {group.comments: group for group in semantic.skip_groups}
        self.assertEqual(3, len(by_comment["Internal ticket"].member_ids))
        self.assertEqual(1, len(by_comment["Other condition"].member_ids))

    def test_lifecycle_duplicates_share_one_record_identity(self) -> None:
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["community_tests"]
        sheet.append([cell.value for cell in sheet[2]])
        workbook.save(self.path)
        workbook.close()

        snapshot = WorkbookLoader().load(self.path, VERSION)
        failure = snapshot.failures[0]
        self.assertEqual(2, len(snapshot.record_excel_rows[failure.record_id]))


if __name__ == "__main__":
    unittest.main()
