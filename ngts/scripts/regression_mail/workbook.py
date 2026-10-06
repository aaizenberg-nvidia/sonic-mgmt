"""Read-only Excel ingestion for collected SONiC results."""

from __future__ import annotations

import json
import os
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, DefaultDict, Dict, List, Mapping, Sequence, Tuple

from ngts.scripts.regression_mail.models import ResultRow, SemanticReport, WorkbookSnapshot
from ngts.scripts.regression_mail.normalization import (
    extract_redmine_urls,
    normalize_text,
    normalize_version,
    stable_record_id,
)


REQUIRED_HEADERS: Sequence[str] = (
    "session_id",
    "mars_key_id",
    "testbed",
    "test name",
    "result",
    "message",
    "topology",
    "hwsku",
    "os_version",
    "sanitized_testname",
)
SUPPORTED_RESULTS = {"pass", "fail", "skipped"}
_RESULT_PRIORITY = {"pass": 0, "skipped": 1, "fail": 2}
FALLBACK_SKIP_REASON = "Internal issue, WIP"
_REMOVE_FAILED_TESTCASE = "REMOVE_FAILED_TESTCASE"


class WorkbookLoader:
    """Load the unique result sheet without mutating the source workbook."""

    def load(self, path: Path, version: str) -> WorkbookSnapshot:
        try:
            import openpyxl
        except ImportError as error:
            raise RuntimeError(
                "openpyxl is required; install ngts/scripts/regression_mail/requirements.txt"
            ) from error

        source = path.resolve()
        if source.suffix.lower() != ".xlsx":
            raise ValueError("Excel input must use the .xlsx extension")
        if not source.is_file():
            raise FileNotFoundError("Excel input does not exist: {}".format(source))

        workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
        try:
            sheet_name, header_row, headers = self._find_result_sheet(workbook)
            (
                selected_rows,
                record_excel_rows,
                setup_platforms,
                setup_testbeds,
                versions,
            ) = self._read_selected_rows(
                workbook[sheet_name],
                header_row,
                headers,
                normalize_version(version),
            )
        finally:
            workbook.close()

        if not selected_rows:
            raise ValueError("no Excel rows match version {!r}".format(version))

        counts = Counter(row.result for row in selected_rows)
        hardware_pairs = sorted(
            {(row.hwsku, row.topology) for row in selected_rows if row.hwsku or row.topology}
        )
        return WorkbookSnapshot(
            source_path=source,
            sheet_name=sheet_name,
            header_row=header_row,
            selected_rows=selected_rows,
            result_counts={name: counts.get(name, 0) for name in sorted(SUPPORTED_RESULTS)},
            hardware_pairs=[tuple(pair) for pair in hardware_pairs],
            versions=versions,
            setup_platforms=setup_platforms,
            setup_testbeds=setup_testbeds,
            record_excel_rows=record_excel_rows,
        )

    @staticmethod
    def _find_result_sheet(workbook: Any) -> Tuple[str, int, List[str]]:
        candidates: List[Tuple[str, int, List[str]]] = []
        required = set(REQUIRED_HEADERS)
        for worksheet in workbook.worksheets:
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True),
                start=1,
            ):
                headers = [str(value).strip() if value is not None else "" for value in row]
                if required.issubset(set(headers)):
                    semantic_headers = [
                        normalize_text(header).lower().replace("_", " ")
                        for header in headers
                        if header
                    ]
                    if len(semantic_headers) != len(set(semantic_headers)):
                        raise ValueError(
                            "result sheet {!r} has duplicate column names".format(worksheet.title)
                        )
                    candidates.append((worksheet.title, row_number, headers))
                    break

        if not candidates:
            raise ValueError(
                "no worksheet contains all required columns: {}".format(", ".join(REQUIRED_HEADERS))
            )
        if len(candidates) != 1:
            raise ValueError(
                "multiple worksheets contain the required result columns: {}".format(
                    ", ".join(candidate[0] for candidate in candidates)
                )
            )
        return candidates[0]

    @staticmethod
    def _read_selected_rows(
        worksheet: Any,
        header_row: int,
        headers: Sequence[str],
        wanted_version: str,
    ) -> Tuple[
        List[ResultRow],
        Dict[str, List[int]],
        Dict[str, str],
        Dict[str, str],
        List[str],
    ]:
        header_index = {header: index for index, header in enumerate(headers) if header}
        deduplicated: Dict[str, ResultRow] = {}
        record_excel_rows: Dict[str, List[int]] = {}
        setup_variants: DefaultDict[str, List[str]] = defaultdict(list)
        platforms_by_setup: DefaultDict[str, List[str]] = defaultdict(list)
        versions = set()
        for excel_row, cells in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            values = {
                header: cells[index] if index < len(cells) else None
                for header, index in header_index.items()
            }
            if normalize_text(values.get("modify_xml")).upper() == _REMOVE_FAILED_TESTCASE:
                continue
            raw_setup = str(values.get("testbed") or "").strip()
            setup = normalize_text(raw_setup)
            platform = normalize_text(values.get("platform"))
            if setup:
                setup_key = setup.casefold()
                for variant in (raw_setup, setup):
                    if variant and variant.casefold() not in {
                        value.casefold()
                        for value in setup_variants[setup_key]
                    }:
                        setup_variants[setup_key].append(variant)
                if platform and platform.casefold() not in {
                    value.casefold()
                    for value in platforms_by_setup[setup_key]
                }:
                    platforms_by_setup[setup_key].append(platform)
            row_version = normalize_version(values.get("os_version"))
            if row_version:
                versions.add(row_version)
            result = normalize_text(values.get("result")).lower()
            if result not in SUPPORTED_RESULTS:
                raise ValueError(
                    "unsupported result {!r} at {}!{}".format(
                        values.get("result"),
                        worksheet.title,
                        excel_row,
                    )
                )

            test_name = normalize_text(values.get("test name"))
            session_id = normalize_text(values.get("session_id"))
            key_id = normalize_text(values.get("mars_key_id"))
            record_id = stable_record_id(row_version, session_id, key_id, test_name)
            record_excel_rows.setdefault(record_id, []).append(excel_row)
            row = ResultRow(
                record_id=record_id,
                excel_row=excel_row,
                session_id=session_id,
                mars_key_id=key_id,
                testbed=normalize_text(values.get("testbed")),
                test_name=test_name,
                sanitized_testname=normalize_text(values.get("sanitized_testname")),
                result=result,
                message=normalize_text(values.get("message")),
                topology=normalize_text(values.get("topology")),
                hwsku=normalize_text(values.get("hwsku")),
                platform=platform,
                values=values,
            )
            existing = deduplicated.get(record_id)
            if existing is None or _RESULT_PRIORITY[row.result] > _RESULT_PRIORITY[existing.result]:
                deduplicated[record_id] = row

        setup_platforms = {}
        for key, variants in setup_variants.items():
            platforms = platforms_by_setup.get(key, [])
            replacement = platforms[0] if len(platforms) == 1 else ""
            for variant in variants:
                setup_platforms[variant] = replacement
        selected_rows = sorted(
            deduplicated.values(),
            key=lambda row: row.excel_row,
        )
        if wanted_version not in versions:
            raise ValueError("no Excel rows match version {!r}".format(wanted_version))
        setup_testbeds: Dict[str, str] = {}
        displays_by_setup: DefaultDict[str, List[str]] = defaultdict(list)
        for row in selected_rows:
            key = row.testbed.casefold()
            display = _display_testbed(
                row.platform,
                row.topology,
                row.testbed,
            )
            if display and display.casefold() not in {
                value.casefold() for value in displays_by_setup[key]
            }:
                displays_by_setup[key].append(display)
        for key, displays in displays_by_setup.items():
            replacement = displays[0] if len(displays) == 1 else ""
            for variant in setup_variants.get(key, ()):
                setup_testbeds[variant] = replacement
        return (
            selected_rows,
            record_excel_rows,
            setup_platforms,
            setup_testbeds,
            sorted(versions),
        )


class ArtifactWriter:
    """Create the generated skip mapping."""

    def write_skips(
        self,
        workbook: WorkbookSnapshot,
        semantic: SemanticReport,
    ) -> Path:
        rows_by_id = {row.record_id: row for row in workbook.skipped}
        urls = set()
        for group in semantic.skip_groups:
            for member in group.member_ids:
                row = rows_by_id.get(member)
                if row is not None:
                    urls.update(extract_redmine_urls(row.message))
            urls.update(url.lower() for url in group.redmine_urls)

        payload = {
            url: FALLBACK_SKIP_REASON
            for url in sorted(urls)
        }
        destination = workbook.source_path.parent / "skips.json"
        _atomic_json(destination, payload)
        return destination


def _display_testbed(
    platform: str,
    topology: str,
    testbed: str,
) -> str:
    platform_name = normalize_text(platform)
    if (
        not platform_name or
        platform_name.casefold() == "unknown platform"
    ):
        return normalize_text(testbed)
    values = [platform_name, normalize_text(topology)]
    return "-".join(value for value in values if value)


def _atomic_json(path: Path, payload: Mapping[str, str]) -> None:
    handle = tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        prefix=".{}.".format(path.name),
        suffix=".tmp",
        dir=str(path.parent),
        delete=False,
    )
    temp_path = Path(handle.name)
    try:
        json.dump(payload, handle, indent=2, ensure_ascii=False, sort_keys=True)
        handle.write("\n")
        handle.flush()
        os.fsync(handle.fileno())
        os.fchmod(handle.fileno(), 0o664)
        handle.close()
        os.replace(temp_path, path)
    except Exception:
        handle.close()
        try:
            temp_path.unlink()
        except FileNotFoundError:
            pass
        raise
